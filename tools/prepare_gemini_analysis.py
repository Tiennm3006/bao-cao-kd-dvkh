#!/usr/bin/env python3
"""Build one Gemini-ready prompt per workbook sheet.

This step deliberately does not call Gemini. It creates an auditable JSON payload
so the API integration can be added without exposing a key in the dashboard.
"""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook


def norm(value: object) -> str:
    text = "" if value is None else str(value)
    text = text.replace("Đ", "D").replace("đ", "d")
    text = unicodedata.normalize("NFD", text)
    text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", "", text.lower())


def compact(value: object) -> str:
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ")
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def used_rows(ws) -> list[list[str]]:
    rows = [[compact(cell) for cell in row] for row in ws.iter_rows(values_only=True)]
    while rows and not any(rows[-1]):
        rows.pop()
    width = max((max((i for i, value in enumerate(row) if value), default=-1) for row in rows), default=-1) + 1
    return [row[:width] + [""] * max(0, width - len(row)) for row in rows]


def markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return "(Sheet không có dữ liệu)"
    width = max(len(row) for row in rows)
    rows = [row + [""] * (width - len(row)) for row in rows]
    # Keep the original worksheet rows. Prompt instructions can interpret merged
    # headers and section rows better than a guessed header transformation.
    return "\n".join(
        "| " + " | ".join(value.replace("|", "\\|") for value in row) + " |"
        for row in rows
    )


def parse_prompt_blocks(path: Path) -> dict[str, str]:
    text = path.read_text(encoding="utf-8")
    matches = list(re.finditer(r"^## Sheet\s+\d+\s+—\s+`([^`]+)`.*$", text, re.MULTILINE))
    blocks: dict[str, str] = {}
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        blocks[match.group(1).strip()] = text[match.start():end].strip()
    if not blocks:
        raise ValueError(f"Không tìm thấy block '## Sheet ...' trong {path}")
    return blocks


def prompt_subject(key: str) -> str:
    """Return a stable subject key independent of changing numeric prefixes."""
    value = norm(key)
    aliases = {
        "1crm": "crm",
        "2matdientbacc": "matdientramlaplai",
        "3hailongkh": "khonghailong",
        "4thutd": "thutien",
        "5nonsnn": "nonsnn",
        "41nonsnn": "nonsnn",
        "6nlmt": "nlmt",
        "5nlmt": "nlmt",
        "7thaydk": "thaydk",
        "6thaydk": "thaydk",
        "8doxa": "doxa",
        "7doxa": "doxa",
        "9truythu": "truythu",
        "8truythu": "truythu",
        "10ktsdd": "ktsdd",
        "9ktsdd": "ktsdd",
        "11dubaophutai": "dubaophutai",
        "9dubaophutai": "dubaophutai",
        "12diennhantuan": "diennhantuan",
        "10diennhantuan": "diennhantuan",
        "13lapdattbdd": "lapdattbdd",
        "11lapdattbdd": "lapdattbdd",
    }
    return aliases.get(value, value)


def build_mapping(prompt_blocks: dict[str, str], sheet_names: Iterable[str]) -> dict[str, tuple[str, str]]:
    by_subject = {prompt_subject(key): (key, block) for key, block in prompt_blocks.items()}
    mapping: dict[str, tuple[str, str]] = {}
    for sheet_name in sheet_names:
        subject = prompt_subject(sheet_name.strip())
        if subject in by_subject:
            mapping[sheet_name] = by_subject[subject]
    return mapping


def inject_prompt(template: str, table: str, previous: str) -> str:
    prompt = re.sub(r"\[DÁN DỮ LIỆU[^\]]*\]", table, template, count=1)
    prompt = re.sub(r"\[DÁN SỐ LIỆU TUẦN TRƯỚC[^\]]*\]", previous or "(Chưa cung cấp dữ liệu tuần trước)", prompt, count=1)
    return prompt


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--previous", type=Path)
    parser.add_argument("--prompts", type=Path, default=Path("Prompt-danh-gia-theo-sheet.md"))
    parser.add_argument("--output", type=Path, default=Path("gemini_analysis_requests.json"))
    args = parser.parse_args()

    prompt_blocks = parse_prompt_blocks(args.prompts)
    workbook = load_workbook(args.workbook, read_only=True, data_only=True)
    previous_wb = load_workbook(args.previous, read_only=True, data_only=True) if args.previous else None
    mapping = build_mapping(prompt_blocks, workbook.sheetnames)
    missing = [name for name in workbook.sheetnames if name not in mapping]

    requests = []
    for sheet_name in workbook.sheetnames:
        if sheet_name not in mapping:
            continue
        prompt_id, template = mapping[sheet_name]
        rows = used_rows(workbook[sheet_name])
        table = markdown_table(rows)
        previous_table = ""
        if previous_wb:
            previous_name = next((n for n in previous_wb.sheetnames if prompt_subject(n.strip()) == prompt_subject(sheet_name.strip())), None)
            if previous_name:
                previous_table = markdown_table(used_rows(previous_wb[previous_name]))
        requests.append({
            "sheet_name": sheet_name,
            "prompt_id": prompt_id,
            "row_count": len(rows),
            "column_count": max((len(row) for row in rows), default=0),
            "prompt": inject_prompt(template, table, previous_table),
        })

    payload = {
        "source_file": str(args.workbook),
        "previous_file": str(args.previous) if args.previous else None,
        "prompt_file": str(args.prompts),
        "sheet_count": len(workbook.sheetnames),
        "mapped_count": len(requests),
        "unmapped_sheets": missing,
        "requests": requests,
    }
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(args.output), "mapped": len(requests), "unmapped": missing}, ensure_ascii=True))


if __name__ == "__main__":
    main()
